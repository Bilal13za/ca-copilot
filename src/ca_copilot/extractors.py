"""Safe text extraction from supported document types.

Design rules:
- Never execute macros or embedded code; we only read text/structure.
- Malformed files raise ExtractionError with a human-readable message.
- Scanned (image-only) PDFs are REPORTED as such; OCR is optional and
  explicit — never claim an image-only PDF was fully read without OCR.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {"pdf", "docx", "txt", "md", "csv", "xlsx"}


class ExtractionError(Exception):
    """Raised when a file cannot be read safely or is unsupported."""


@dataclass
class PageText:
    page: int | None          # 1-based page number, or None if n/a
    section: str | None       # sheet name / heading context, or None
    text: str
    image_only: bool = False  # True for scanned pages with no extractable text


@dataclass
class Extraction:
    file_type: str
    pages: list[PageText] = field(default_factory=list)
    page_count: int | None = None
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def has_extractable_text(self) -> bool:
        return any(p.text.strip() for p in self.pages)


def detect_type(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext not in SUPPORTED:
        raise ExtractionError(
            f"Unsupported file type '.{ext}'. Supported: PDF, DOCX, TXT, Markdown, CSV, XLSX."
        )
    return ext


def extract(path: Path, *, ocr_enabled: bool = False, tesseract_cmd: str = "",
            ocr_dpi: int = 200) -> Extraction:
    """Extract text from a supported file. Never executes embedded code."""
    file_type = detect_type(path)
    try:
        if file_type == "pdf":
            return _extract_pdf(path, ocr_enabled, tesseract_cmd, ocr_dpi)
        if file_type == "docx":
            return _extract_docx(path)
        if file_type in ("txt", "md"):
            return _extract_text(path)
        if file_type == "csv":
            return _extract_csv(path)
        if file_type == "xlsx":
            return _extract_xlsx(path)
    except ExtractionError:
        raise
    except Exception as exc:  # malformed file -> friendly error, never a traceback in UI
        raise ExtractionError(f"Could not read '{path.name}': {exc}") from exc
    raise ExtractionError(f"Unsupported file type '.{file_type}'.")


# ---------------------------------------------------------------- PDF

def _extract_pdf(path: Path, ocr_enabled: bool, tesseract_cmd: str, dpi: int) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF support needs the 'pypdf' package.") from exc
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ExtractionError("This PDF is encrypted/password-protected and cannot be read.")
    pages: list[PageText] = []
    scanned_pages = 0
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            scanned_pages += 1
            pages.append(PageText(page=i, section=None, text="", image_only=True))
        else:
            pages.append(PageText(page=i, section=None, text=text))
    ext = Extraction(file_type="pdf", pages=pages, page_count=len(reader.pages))
    if scanned_pages and not ocr_enabled:
        ext.warnings.append(
            f"{scanned_pages} of {len(reader.pages)} page(s) appear to be scanned images "
            "with no extractable text. Enable OCR in Settings to attempt reading them."
        )
    elif scanned_pages and ocr_enabled:
        _ocr_pdf_pages(path, ext, tesseract_cmd=tesseract_cmd, dpi=dpi)
    return ext


def _ocr_pdf_pages(path: Path, ext: Extraction, *, tesseract_cmd: str, dpi: int) -> None:
    """Optional OCR pass over image-only pages. Explicit opt-in only."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        ext.warnings.append("OCR requested but 'pytesseract'/'Pillow' are not installed.")
        raise ExtractionError("OCR needs 'pytesseract' and 'Pillow' installed.") from exc
    try:
        from pypdf import PdfReader  # already imported above, but keep local for clarity
        _ = PdfReader
        # Render pages via pdf2image-like approach is heavy; instead use embedded images.
        reader_pages = ext.pages
        for pt in reader_pages:
            if not pt.image_only:
                continue
            # Extract embedded images from the page and OCR them.
            imgs = _page_images(path, pt.page or 1)
            ocr_texts = []
            for img in imgs:
                if tesseract_cmd:
                    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
                ocr_texts.append(pytesseract.image_to_string(img))
            combined = "\n".join(t.strip() for t in ocr_texts if t.strip())
            if combined.strip():
                pt.text = combined
                pt.image_only = False
                ext.ocr_used = True
        remaining = sum(1 for p in ext.pages if p.image_only)
        if remaining:
            ext.warnings.append(
                f"{remaining} page(s) still have no readable text after OCR."
            )
    except ExtractionError:
        raise
    except Exception as exc:
        ext.warnings.append(f"OCR failed on some pages: {exc}")


def _page_images(path: Path, page_number: int):
    """Return PIL images embedded in a PDF page (no rendering of vector pages)."""
    from pypdf import PdfReader
    from PIL import Image
    import io
    reader = PdfReader(str(path))
    page = reader.pages[page_number - 1]
    images = []
    try:
        for img in page.images:
            images.append(Image.open(io.BytesIO(img.data)))
    except Exception:
        pass
    return images


# ---------------------------------------------------------------- DOCX

def _extract_docx(path: Path) -> Extraction:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("DOCX support needs the 'python-docx' package.") from exc
    doc = docx.Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:  # tables are common in study notes
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    pages = [PageText(page=None, section=None, text=text)] if text.strip() else []
    return Extraction(file_type="docx", pages=pages)


# ---------------------------------------------------------------- TXT / Markdown

def _extract_text(path: Path) -> Extraction:
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        raise ExtractionError("Could not decode text file (tried UTF-8 and Windows-1252).")
    pages = [PageText(page=None, section=None, text=text)] if text.strip() else []
    return Extraction(file_type="txt", pages=pages)


# ---------------------------------------------------------------- CSV

def _extract_csv(path: Path) -> Extraction:
    rows: list[list[str]] = []
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        try:
            sample = fh.read(4096)
            fh.seek(0)
            dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(fh, dialect)
        for i, row in enumerate(reader):
            if i > 20000:  # cap indexing; the workbench reads full data separately
                break
            rows.append(row)
    lines = [", ".join(c for c in r) for r in rows]
    text = "\n".join(lines)
    pages = [PageText(page=None, section="csv", text=text)] if text.strip() else []
    ext = Extraction(file_type="csv", pages=pages)
    if len(rows) > 20000:
        ext.warnings.append("Only the first 20,000 CSV rows were indexed for search.")
    return ext


# ---------------------------------------------------------------- XLSX

def _extract_xlsx(path: Path) -> Extraction:
    try:
        import openpyxl
    except ImportError as exc:
        raise ExtractionError("XLSX support needs the 'openpyxl' package.") from exc
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    pages: list[PageText] = []
    try:
        for ws in wb.worksheets:
            lines: list[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 20000:
                    lines.append("… (sheet truncated for search indexing)")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    lines.append(" | ".join(cells))
            text = "\n".join(lines)
            if text.strip():
                # Sheet name goes into the indexed text so search finds it,
                # and into `section` for precise citations.
                pages.append(PageText(page=None, section=ws.title,
                                      text=f"Sheet: {ws.title}\n{text}"))
    finally:
        wb.close()
    return Extraction(file_type="xlsx", pages=pages, page_count=len(wb.sheetnames))


# ---------------------------------------------------------------- Chunking

def chunk_text(text: str, *, page: int | None, section: str | None,
               chunk_chars: int = 1500, overlap_chars: int = 200) -> list[tuple[int | None, str | None, str]]:
    """Split text into overlapping chunks for retrieval. Returns (page, section, text)."""
    text = text.strip()
    if not text:
        return []
    chunks: list[tuple[int | None, str | None, str]] = []
    step = max(1, chunk_chars - overlap_chars)
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_chars)
        # Prefer to break at a paragraph/line boundary near the end.
        if end < len(text):
            cut = text.rfind("\n", start, end)
            if cut > start + chunk_chars // 2:
                end = cut
        piece = text[start:end].strip()
        if piece:
            chunks.append((page, section, piece))
        if end >= len(text):
            break
        start = end - overlap_chars
    return chunks
